# -*- coding: utf-8 -*-
"""把 Markdown 指南转换为 Word 文档(.docx)"""
import re
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

SRC = r"D:\AI\Claude Project\招投标采购Agent项目_从零搭建指南.md"
DST = r"D:\AI\Claude Project\招投标采购Agent项目_从零搭建指南.docx"

doc = Document()

# ---------- 全局样式:中文字体 ----------
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(10.5)
style.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

for hname, size, color in [("Heading 1", 22, "1F4E79"), ("Heading 2", 16, "2E74B5"),
                           ("Heading 3", 13, "2E74B5"), ("Heading 4", 11, "404040")]:
    h = doc.styles[hname]
    h.font.name = "Calibri"
    h.font.size = Pt(size)
    h.font.bold = True
    h.font.color.rgb = RGBColor.from_string(color)
    h.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")


def set_cn(run):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")


def add_inline(par, text):
    """解析行内 **粗体** 与 `代码`"""
    for tok in re.split(r"(\*\*.*?\*\*|`[^`]*`)", text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            r = par.add_run(tok[2:-2]); r.bold = True; set_cn(r)
        elif tok.startswith("`") and tok.endswith("`"):
            r = par.add_run(tok[1:-1]); r.font.name = "Consolas"
            r.font.size = Pt(9.5); r.font.color.rgb = RGBColor.from_string("C7254E")
            r._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        else:
            r = par.add_run(tok); set_cn(r)


def add_code_block(code_lines):
    for line in code_lines:
        par = doc.add_paragraph()
        par.paragraph_format.space_before = Pt(0)
        par.paragraph_format.space_after = Pt(0)
        par.paragraph_format.left_indent = Inches(0.2)
        r = par.add_run(line if line else " ")
        r.font.name = "Consolas"; r.font.size = Pt(9)
        r.font.color.rgb = RGBColor.from_string("24292E")
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        # 浅灰底纹
        par._p.get_or_add_pPr().append(
            par._p.get_or_add_pPr().makeelement(qn("w:shd"),
                {qn("w:val"): "clear", qn("w:fill"): "F2F2F2"}))


def add_table(rows):
    ncol = max(len(r) for r in rows)
    t = doc.add_table(rows=0, cols=ncol)
    t.style = "Light Grid Accent 1"
    for i, row in enumerate(rows):
        cells = t.add_row().cells
        for j in range(ncol):
            txt = row[j].strip() if j < len(row) else ""
            p = cells[j].paragraphs[0]
            add_inline(p, txt)
            for r in p.runs:
                r.font.size = Pt(9)
                if i == 0:
                    r.bold = True


lines = open(SRC, encoding="utf-8").read().splitlines()
i = 0
n = len(lines)
while i < n:
    line = lines[i]
    s = line.strip()

    # 代码块
    if s.startswith("```"):
        j = i + 1
        while j < n and not lines[j].strip().startswith("```"):
            j += 1
        add_code_block(lines[i + 1:j])
        i = j + 1
        continue

    # 标题
    m = re.match(r"^(#{1,4})\s+(.*)$", s)
    if m:
        doc.add_heading(m.group(2).strip(), level=len(m.group(1)))
        i += 1
        continue

    # 分隔线
    if re.match(r"^-{3,}$", s):
        i += 1
        continue

    # 表格(连续以 | 开头的行)
    if s.startswith("|"):
        rows = []
        while i < n and lines[i].strip().startswith("|"):
            cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
            if not all(re.match(r"^:?-{2,}:?$", c) for c in cells if c):
                rows.append(cells)
            i += 1
        if rows:
            add_table(rows)
        continue

    # 引用块
    if s.startswith(">"):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        add_inline(p, s.lstrip("> ").strip())
        for r in p.runs:
            r.italic = True
            r.font.color.rgb = RGBColor.from_string("595959")
        i += 1
        continue

    # 列表
    if s:
        m = re.match(r"^([-*]|\d+\.)\s+(.*)$", s)
        if m:
            ordered = m.group(1)[0].isdigit()
            p = doc.add_paragraph(style="List Number" if ordered else "List Bullet")
            add_inline(p, m.group(2))
            i += 1
            continue

    # 空行
    if not s:
        i += 1
        continue

    # 普通段落
    p = doc.add_paragraph()
    add_inline(p, s)
    i += 1

doc.save(DST)
print("OK:", DST)
